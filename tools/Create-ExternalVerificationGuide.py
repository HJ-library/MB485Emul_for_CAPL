from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "05_Test" / "External_Verification_Guide.docx"
RELEASE_VERSION = "0.5.0"

NAVY = "17365D"
BLUE = "2F75B5"
PALE_BLUE = "D9EAF7"
PALE_GRAY = "F2F4F7"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "FCE4D6"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_text(cell, text, bold=False, color="000000", size=8.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles[f"Heading {level}"]
    run = paragraph.add_run(text)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_test_table(doc, rows):
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Cm(1.5), Cm(2.4), Cm(6.2), Cm(6.4), Cm(2.2)]
    headers = ["Test ID", "Phase", "시험 내용", "기대 결과", "증적"]
    for idx, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        set_cell_text(cell, header, True, "FFFFFF", 8.5)
        shade(cell, NAVY)
        cell.width = widths[idx]
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            set_cell_text(cells[idx], value, idx == 0, "000000", 8.0)
            cells[idx].width = widths[idx]
            if row_index % 2 == 1:
                shade(cells[idx], PALE_GRAY)
    doc.add_paragraph()


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(8.6)
    normal.paragraph_format.space_after = Pt(2.5)
    normal.paragraph_format.line_spacing = 1.0
    for name, size, color in [("Heading 1", 15, NAVY), ("Heading 2", 11, BLUE)]:
        style = styles[name]
        style.font.name = "Malgun Gothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.space_after = Pt(3)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("MODBUS485EMULATOR  |  EXTERNAL VERIFICATION")
    r.font.name = "Aptos"
    r.font.size = Pt(8)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Modbus485Emulator\n외부 실환경 검증 가이드")
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    sub = doc.add_paragraph("CANalyzer 9.x SP8 x86 · CANalyzer 18 x64 · SystemBase Multi USB 485 · 실제 Board")
    sub.paragraph_format.space_after = Pt(10)
    sub.runs[0].font.size = Pt(10)
    sub.runs[0].font.color.rgb = RGBColor.from_string(BLUE)

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    for row, (label, value) in enumerate([
        ("문서 목적", "자동 검증으로 대체할 수 없는 대상 실환경의 CAPL, USB-RS485 및 Board 시험 절차를 정의한다."),
        ("현재 상태", f"Release {RELEASE_VERSION} Baseline. Windows x86/x64 자동 Build와 ABI 검증 완료. CANalyzer 9.x SP8 및 CANalyzer 18 Measurement와 장비 항목은 External Verification Pending"),
        ("결과 기록", "Modbus485Emulator_External_Test_Result.xlsx의 Result를 Not Tested에서 갱신한다."),
        ("주의", "가상 CAN 채널만 사용하며 물리 CAN Bus Load를 합격 기준으로 사용하지 않는다."),
    ]):
        set_cell_text(table.cell(row, 0), label, True, "FFFFFF")
        shade(table.cell(row, 0), NAVY)
        set_cell_text(table.cell(row, 1), value)
        shade(table.cell(row, 1), PALE_BLUE if row < 3 else AMBER)

    add_heading(doc, "1. 시험 전 준비", 1)
    add_bullets(doc, [
        "Release/Win32와 Release/x64 M485E.dll, M485E.vmodule, CAPL 파일, M485E.vsysvar, DeviceMap/FrameRule/참조 Register Map을 동일 Baseline에서 준비한다.",
        "CANalyzer 9.x SP8 x86 또는 CANalyzer 18 x64 Configuration에 활성 Port 수(1~4)만큼 가상 CAN 채널을 생성하고 Main의 PortSlot별 채널 번호와 맞춘다.",
        "SystemBase Multi USB 485의 시험 대상 활성 Port를 Board RS-485에 연결한다. Adapter의 자동 송수신 방향 제어를 사용하며 DLL은 RTS 방향 제어를 수행하지 않는다.",
        "Board Modbus Master의 Response Timeout, 요청 주기, 실제 COM 번호와 Serial 설정을 결과 파일 Environment 시트에 기록한다.",
        "시험 시작 전 Git Commit, DLL SHA-256, CANalyzer 버전/SP, Vector Driver 버전을 기록한다.",
    ])

    add_heading(doc, "2. CANalyzer 구성 절차", 1)
    steps = [
        "tools/Prepare-M485EPackage.ps1 -Configuration Release를 실행해 DLL과 생성 계약을 Config/CAPL에 배치한다.",
        "CAPL Program Node에 Config/CAPL/M485E/M485E_Main.can을 연결한다. Main은 생성 계약, Core, Monitoring, UI를 포함한다.",
        "Config/System Variables/M485E.vsysvar를 Configuration에 추가한다.",
        "M485E::Monitoring::Btn_ReadOption을 int32 Scalar 편의 Button에 연결한다. 이 변수에는 별도 CAPL Handler가 없으며 Check_ReadOption[200], 100ms Candidate 동기화와 1,000ms Cycle 적용 정책을 그대로 사용한다.",
        "Panel Designer에 CAPL Output View Control을 추가하고 Output Mode를 Overwrite로 설정한다. Main.can 설정값의 Panel Name 'M485E'와 Control Name 'RtuLogOutput'을 실제 Panel/Control 이름과 정확히 일치시킨다.",
        "Main.can에서 DeviceMap과 FrameRule 상대경로를 지정한다. Core가 getAbsFilePath()로 절대경로를 DLL에 전달하고 Register Map은 DeviceMap의 CSV_DB_Name으로 선택한다.",
        "gM485E_ActivePortCount를 DeviceMap 고유 Port 수와 같은 1~4로 설정하고 앞쪽 PortSlot의 COM 번호와 BaudRate/DataBits/Parity/StopBits를 설정한다. 기본 배포 Source Contract가 두 수의 일치를 검사하며, Refresh 완료를 Polling한 뒤 Connect한다.",
        "ReadOption 정적 매핑을 설정하고 ReadSwitch ON에서 Monitoring을 시작한다.",
    ]
    for idx, step in enumerate(steps, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.25)
        p.add_run(f"{idx}. ").bold = True
        p.add_run(step)

    add_heading(doc, "3. 필수 시험 순서", 1)
    rows = [
        ("EV-001", "1/2", "CANalyzer 9.x x86 직접 DLL 및 CANalyzer 18 x64 .vmodule CAPL Compile/로드", "두 환경에서 23개 API가 중복/누락 없이 Compile되고 Measurement가 시작됨", "Compile log"),
        ("EV-002", "3/8", "DeviceMap/FrameRule 두 상대경로와 참조 Register Map 준비 후 비동기 Refresh 상태 Polling", "Disconnected, ConfigurationValid=1, LastResult=0", "Write/Status"),
        ("EV-003", "7/8", "1~4개 활성 COM Port 전체 Connect", "모든 활성 Port가 열릴 때만 Connected; 일부 실패 시 전체 Rollback", "Status/handle"),
        ("EV-004", "5/7", "Port별 Slave에 FC03 Quantity 1/125 요청", "직접 Address의 Holding 값, CRC Low/High 정상 응답", "Serial trace"),
        ("EV-005", "5/7", "FC04 정상/권한 공백 요청", "정상값 또는 Exception 0x02, 부분 응답 없음", "Serial trace"),
        ("EV-006", "4/5/7", "FC10 정상 Write, 동일값 Write, 권한 공백 Write", "원자 적용; 다음 선택된 0x80에 최종값 반영; 공백은 0x02 및 무변경", "Serial/CAN"),
        ("EV-007", "5/7", "CRC 오류, FC06, Quantity/ByteCount 오류", "CRC 오류 무응답; FC06=0x01; 값 오류=0x03", "Serial trace"),
        ("EV-008", "6/9", "Periodic Rule Data Quantity 6", "0x80 Header 1 + Data 2, Header Byte 3=0, 마지막 DLC 4", "CAN trace"),
        ("EV-009", "6/9", "RTU/Direct/Reset 뒤 최종값 확인", "0x81 미출력, 다음 선택된 0x80에 최종값 반영, 중간 상태 재생 없음", "CAN trace"),
        ("EV-010", "6/9", "여러 활성 Port Periodic 출력", "활성 Port Round-robin, Rule Data 묶음 비분할, 0x80만 출력", "CAN trace"),
        ("EV-011", "9/10", "ReadSwitch OFF/ON 및 ReadOption 변경", "OFF 중 RTU 유지·CAN 억제, 다음 Cycle에 Candidate 원자 반영", "CAN/Write"),
        ("EV-012", "10", "Panel Manual Write Data[122] 및 Signal Control", "Send 시 요청 Snapshot 사용, Quantity 1~122와 word 범위 검증, ResultCode 표시", "Panel capture"),
        ("EV-013", "10", "Panel Signal Float/Signed Read/Write", "12.5 등 물리값이 그대로 왕복하고 실패 시 호출 전 출력 유지", "Panel capture"),
        ("EV-014", "7/8", "한 USB-RS485 Port 제거 후 복구", "전체 정리 후 3초 간격 최대 10회; 복구 시 모든 활성 Port Connected", "Timestamp log"),
        ("EV-015", "8", "Disconnect/Connect와 Finalize", "Disconnect는 값 보존, Finalize는 2초 목표 내 Join/Close 후 Context 폐기", "Status/Process"),
        ("EV-016", "9/10", "200 Rule, 최대 Quantity, 66 Frame/Tick 성능 Profile", "RTU 영향 없이 선택 Rule의 0x80을 1000ms Cycle로 처리; Error Queue DropCount 변화/Overflow 최초 전환은 누적 통계 한 번 출력", "Trace/metrics"),
        ("EV-017", "5/7/10", "Index별 RTU 통신 OFF/ON", "비활성 Index는 FC03/04/10에 무응답·무변경, 동일 Port의 다른 Slave와 Direct/CAN 제어는 유지; Refresh 성공 후 활성화", "Panel/Serial/CAN"),
        ("EV-018", "7/10", "RTU Log Switch와 CSV 기록", "OFF→ON마다 새 파일, REQ/RES 또는 REQ/NO_RES 동일 Sequence, 12개 Column·전체 ADU·DroppedCount·ActiveLed 확인", "CSV/Panel"),
        ("EV-019", "10", "Btn_ReadOption과 RTU Log CAPL Output View", "Btn_ReadOption은 int32 Scalar/no-handler이고 기존 Check_ReadOption 정책을 유지한다. LogSwitch ON에서 전체 REQ/RES/NO_RES는 CSV에 보존하며 동일 Drain 문자열의 최근 10줄만 200ms 이상 간격으로 Output View에 Overwrite한다. Panel/Control 이름 불일치가 CSV 기록을 실패시키지 않으며 OFF 후 마지막 표시는 유지된다.", "Panel capture/CSV/timing"),
    ]
    add_test_table(doc, rows)

    add_heading(doc, "4. CAN Trace 판독 기준", 1)
    add_bullets(doc, [
        "실제 Arbitration ID는 0x00IIRRMM(Periodic)이고 Vector 내부 Extended 표시는 0x80IIRRMM이다. 0x81 변경 Message는 출력하지 않는다.",
        "Header DLC는 8이며 Index, Slave ID, FC, 고정값 0x00, Address High/Low, Quantity High/Low 순서다.",
        "Data는 Address 오름차순 Register를 High Byte/Low Byte로 최대 4개씩 저장하며 마지막 Message는 패딩하지 않는다.",
        "가상 CAN이므로 물리 Bus Load, CAN BaudRate, 75/85/90% 자동 정지를 평가하지 않는다. Output Pump 실행시간, Deadline Miss, Queue 상태와 RTU 영향으로 판정한다.",
    ])

    add_heading(doc, "5. 실패 시 수집할 증적", 1)
    add_bullets(doc, [
        "결과 Workbook의 Actual Result, Result, Evidence, Notes를 즉시 갱신하고 원본 Trace/로그 파일은 변경하지 않는다.",
        "CSV 오류는 FileType, Row, Column, ErrorCode 및 CAPL Write 문장을 함께 저장한다.",
        "Serial 오류는 PortSlot, COM 이름, Stage, Win32 OsError와 발생 시각을 함께 저장한다.",
        "성능 실패는 활성 Rule 목록, Rule별 Quantity, Pump Tick 실행시간, Periodic Pending 및 Board Timeout을 함께 저장한다.",
        "Error Queue DropCount 또는 Overflow Latch가 변하면 Write Window의 Pending, Capacity, High-water Mark, DropCount 및 Overflow 누적 통계를 저장한다.",
    ])

    add_heading(doc, "6. 종료 및 판정", 1)
    p = doc.add_paragraph()
    p.add_run("PASS 조건: ").bold = True
    p.add_run("해당 Test ID의 Expected Result를 충족하고 Evidence가 연결되어야 한다. ")
    p.add_run("미실행 항목은 반드시 Not Tested로 유지하며 PASS로 추정하지 않는다.")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Modbus485Emulator · External Verification Guide")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("6B7280")

    doc.core_properties.title = "Modbus485Emulator External Verification Guide"
    doc.core_properties.subject = f"Release {RELEASE_VERSION} external verification procedures"
    doc.core_properties.author = "Modbus485Emulator Project"
    doc.core_properties.last_modified_by = "Modbus485Emulator Project"
    doc.core_properties.comments = ""
    doc.core_properties.category = "Software Engineering"
    doc.core_properties.keywords = "Modbus485Emulator, M485E, Modbus RTU, CAPL"
    doc.core_properties.created = datetime(2026, 8, 12)
    doc.core_properties.modified = datetime(2026, 8, 17)
    doc.core_properties.version = RELEASE_VERSION
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
