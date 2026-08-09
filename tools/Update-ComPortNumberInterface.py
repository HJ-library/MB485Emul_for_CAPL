from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SRS = ROOT / "docs" / "01_Requirements" / "Modbus485Emulator_SRS.docx"
FS = ROOT / "docs" / "02_FunctionalSpecification" / "Modbus485Emulator_Functional_Specification.docx"
SDS = ROOT / "docs" / "03_Design" / "Modbus485Emulator_Software_Design.docx"
GUIDE = ROOT / "docs" / "05_Test" / "External_Verification_Guide.docx"


def set_font_in_cell(cell, size=8.5):
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Malgun Gothic"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
            run.font.size = Pt(size)


def set_cell(cell, text):
    cell.text = text
    set_font_in_cell(cell)


def append_history(document, version, status, description):
    history = document.tables[1]
    if any(row.cells[0].text == version for row in history.rows[1:]):
        return
    cells = history.add_row().cells
    values = [version, "2026-08-09", status, description]
    for cell, value in zip(cells, values):
        set_cell(cell, value)


def update_metadata(document, version):
    metadata = document.tables[0]
    set_cell(metadata.rows[3].cells[1], f"{version} / Baseline Candidate")


def update_srs():
    document = Document(SRS)
    update_metadata(document, "0.10")
    append_history(
        document,
        "0.10",
        "Baseline Candidate",
        "CAPL COM Port 입력을 숫자 ComPortNumber로 변경하고 저수준 DLL 문자열 ABI 유지",
    )
    api_table = document.tables[17]
    set_cell(
        api_table.rows[4].cells[1],
        "M485E_Core_SetPortConfig(PortSlot, ComPortNumber, BaudRate, DataBits, Parity, StopBits)",
    )
    temporary = SRS.with_suffix(".updated.docx")
    document.save(temporary)
    temporary.replace(SRS)


def update_functional_specification():
    document = Document(FS)
    update_metadata(document, "0.11")
    append_history(
        document,
        "0.11",
        "Baseline Candidate",
        "CAPL 숫자 COM Port 입력과 Core의 COM%d 문자열 변환 동작 반영",
    )

    initialization = document.tables[5]
    set_cell(
        initialization.rows[3].cells[1],
        "세 상대경로, PortSlot 0/1, ComPortNumber(1~999), BaudRate, DataBits, Parity, StopBits",
    )
    set_cell(
        initialization.rows[5].cells[1],
        "1. Main.can이 프로젝트 상대경로와 숫자 COM Port 번호를 보관한다.\n"
        "2. getAbsFilePath()로 각 경로를 절대경로로 변환한다.\n"
        "3. Core가 ComPortNumber를 COM%d 문자열로 변환한다.\n"
        "4. Core가 FileType별 경로와 Port 설정을 DLL에 전달한다.\n"
        "5. DLL은 설정만 저장하고 COM Port를 열지 않는다.",
    )

    port_configuration = document.tables[13]
    set_cell(
        port_configuration.rows[3].cells[1],
        "PortSlot, ComPortNumber(1~999), BaudRate, DataBits, Parity, StopBits",
    )
    set_cell(
        port_configuration.rows[5].cells[1],
        "1. ComPortNumber와 나머지 입력 범위를 검증한다.\n"
        "2. CAPL Core가 ComPortNumber를 COM%d 문자열로 변환한다.\n"
        "3. 변환한 Port를 DeviceMap의 Port와 매칭한다.\n"
        "4. Connect 중 DCB/Timeout을 적용한다.\n"
        "5. RTS/방향 제어는 수행하지 않는다.",
    )
    temporary = FS.with_suffix(".updated.docx")
    document.save(temporary)
    temporary.replace(FS)


def update_design():
    document = Document(SDS)
    update_metadata(document, "0.15")
    append_history(
        document,
        "0.15",
        "Implementation Baseline",
        "CAPL Port 설정 구조체 배열과 기존 DLL 스칼라 ABI 경계 반영",
    )
    marker = "28. Numeric COM Port Adapter"
    if not any(paragraph.text == marker for paragraph in document.paragraphs):
        document.add_heading(marker, level=1)
        document.add_paragraph(
            "Main.can은 PortSlot별 COM 번호를 CAPL long으로 관리한다. "
            "M485E_Core_SetPortConfig는 ComPortNumber 1~999를 검증한 뒤 snprintf로 COM%d 문자열을 "
            "생성하여 기존 M485E_Dll_SetPortConfig 문자열 ABI에 전달한다. DeviceMap.csv의 Port 칼럼과 "
            "Win32 Serial 계층은 COM11 같은 Windows Port 이름을 계속 사용한다."
        )
        document.add_paragraph(
            "이 Adapter 경계로 CAPL 사용자는 숫자만 설정하며, Phase 1에서 확정한 21개 DLL Export의 "
            "파라미터와 ABI는 변경하지 않는다."
        )
    structure_marker = "29. CAPL Port Configuration Table"
    if not any(paragraph.text == structure_marker for paragraph in document.paragraphs):
        document.add_heading(structure_marker, level=1)
        document.add_paragraph(
            "Main.can은 두 Port의 portSlot, comPortNumber, baudRate, dataBits, parity, stopBits를 "
            "M485E_PortConfig[2] 구조체 배열로 보관한다. 측정 전 초기화 함수가 각 원소를 채우고 "
            "적용 함수가 elcount() 범위로 순회하여 M485E_Core_SetPortConfig를 호출한다."
        )
        document.add_paragraph(
            "M485E_PortConfig는 CAPL Main 내부 설정 모델이며 DLL 경계를 통과하지 않는다. Core와 DLL은 "
            "기존 스칼라 파라미터 호출을 유지하므로 CAPL 구조체의 메모리 배치가 공개 ABI에 영향을 주지 않는다."
        )
    temporary = SDS.with_suffix(".updated.docx")
    document.save(temporary)
    temporary.replace(SDS)


def update_external_guide():
    document = Document(GUIDE)
    for paragraph in document.paragraphs:
        if "두 PortSlot의 COM 이름과 BaudRate/DataBits/Parity/StopBits" in paragraph.text:
            paragraph.text = paragraph.text.replace(
                "두 PortSlot의 COM 이름과 BaudRate/DataBits/Parity/StopBits",
                "두 PortSlot의 숫자 COM Port 번호(예: 11, 12)와 BaudRate/DataBits/Parity/StopBits",
            )
            for run in paragraph.runs:
                run.font.name = "Malgun Gothic"
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
                run.font.size = Pt(8.6)
    temporary = GUIDE.with_suffix(".updated.docx")
    document.save(temporary)
    temporary.replace(GUIDE)


if __name__ == "__main__":
    update_srs()
    update_functional_specification()
    update_design()
    update_external_guide()
    print(SRS)
    print(FS)
    print(SDS)
    print(GUIDE)
